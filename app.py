import os
import copy
import base64
import io
import subprocess
import tempfile
import uuid
from datetime import datetime

import fitz  # PyMuPDF
import requests

from flask import Flask, render_template, request, send_file, jsonify
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph as DocxParagraph

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 300 * 1024 * 1024  # 300 MB

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), 'WorkingTemplate.docx')
REPORTS_DIR   = os.path.join(os.path.dirname(__file__), 'reports')
UPLOADS_DIR   = '/tmp/report_img_uploads'
os.makedirs(REPORTS_DIR, exist_ok=True)
os.makedirs(UPLOADS_DIR, exist_ok=True)

CHECKLIST_ITEMS = [
    "Choose elevator(s) to be used (service/resident)",
    "Test connectivity at elevator (Wi-Fi / LTE)",
    "Confirm elevator compatibility (door/space)",
    "Identify docking location(s)",
    "Confirm power availability at docking",
    "Confirm access doors or obstacles along robot route",
]

OWNERS = ["Building", "Client", "TII", "Vendor", "TBD"]

TELEGRAM_BOT_TOKEN = (
    os.environ.get('TELEGRAM_BOT_TOKEN')
    or os.environ.get('TELEGRAM_TOKEN')
    or os.environ.get('telegram_token')
    or ''
).strip()
TELEGRAM_CHAT_ID = (
    os.environ.get('TELEGRAM_CHAT_ID')
    or os.environ.get('telegram_chat_id')
    or ''
).strip()


# ── Checkbox formatters ────────────────────────────────────────────────────

def fmt_check(done):
    if done == 'Yes':     return 'Yes ☑'
    if done == 'No':      return 'No ☐'
    if done == 'Pending': return 'Pending ◷'
    return ''

def fmt_elev_type(val):
    return '  '.join(('☑' if o == val else '☐') + ' ' + o for o in ['Service', 'Resident'])

def fmt_connectivity(val):
    return '  '.join(('☑' if o == val else '☐') + ' ' + o for o in ['Good', 'Weak', 'No signal'])

def fmt_owners(selected):
    return '  '.join(('☑' if o in selected else '☐') + ' ' + o for o in OWNERS)

def fmt_pilot(val):
    return '  '.join(('☑' if o == val else '☐') + ' ' + o for o in ['Yes', 'No', 'TBD'])


# ── Placeholder replacement helpers ───────────────────────────────────────

def _replace_in_element(element, replacements):
    """
    Replace <<Key>> markers in all w:p elements within `element`.
    Merges split runs per paragraph before replacing.
    """
    for para in element.iter(qn('w:p')):
        runs = para.findall('.//' + qn('w:r'))
        if not runs:
            continue
        texts = []
        for r in runs:
            t = r.find(qn('w:t'))
            texts.append(t.text or '' if t is not None else '')
        full_text = ''.join(texts)
        if not any(f'<<{k}>>' in full_text for k in replacements):
            continue
        first_run = runs[0]
        first_t = first_run.find(qn('w:t'))
        if first_t is None:
            continue
        first_t.text = full_text
        if full_text.startswith(' ') or full_text.endswith(' '):
            first_t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        for run in runs[1:]:
            parent = run.getparent()
            if parent is not None:
                parent.remove(run)
        for key, value in replacements.items():
            first_t.text = (first_t.text or '').replace(f'<<{key}>>', value)


def _set_cell_text(cell_el, text):
    """Replace all content in a table cell with a single plain-text paragraph."""
    for p in list(cell_el.findall(qn('w:p'))):
        cell_el.remove(p)
    p = OxmlElement('w:p')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = text or ''
    if t.text.startswith(' ') or t.text.endswith(' '):
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    p.append(r)
    cell_el.append(p)


def _add_extra_changes(changes_tbl, changes, has_pilot=False, template_slots=2):
    """
    Append rows beyond the template's pre-filled slots.
    Both D1 and D2 tables: Change | Owner | Notes
    D1 has 2 template slots; D2 has 3 template slots.
    """
    if len(changes) <= template_slots:
        return
    data_rows = changes_tbl.findall(qn('w:tr'))[1:]
    if not data_rows:
        return
    template_row = data_rows[-1]
    for change in changes[template_slots:]:
        new_row = copy.deepcopy(template_row)
        for t_el in new_row.iter(qn('w:t')):
            if '<<' in (t_el.text or ''):
                t_el.text = ''
        cells = new_row.findall(qn('w:tc'))
        _set_cell_text(cells[0], change.get('change', ''))
        _set_cell_text(cells[1], fmt_owners(change.get('owners', [])))
        _set_cell_text(cells[2], change.get('notes', ''))
        template_row.addnext(new_row)
        template_row = new_row


def _remove_empty_change_rows(changes_tbl):
    """Remove data rows from a changes table where both Change and Notes are empty."""
    rows = changes_tbl.findall(qn('w:tr'))
    for row in rows[1:]:   # skip header row
        cells = row.findall(qn('w:tc'))
        if len(cells) >= 3:
            change_text = ''.join(t.text or '' for t in cells[0].iter(qn('w:t'))).strip()
            notes_text  = ''.join(t.text or '' for t in cells[2].iter(qn('w:t'))).strip()
            if not change_text and not notes_text:
                row.getparent().remove(row)


def _insert_images(doc, element, images):
    """
    Replace every '[ Image Placeholder ]' paragraph inside element
    with the corresponding uploaded image.  Images beyond the template's
    4 slots are appended after the last placeholder.
    Returns list of original image bytes (or None) in slot order.
    """
    orig_bytes_list = []
    slot_idx = 0
    last_placeholder_el = None

    for para_el in element.iter(qn('w:p')):
        runs = para_el.findall('.//' + qn('w:r'))
        text = ''.join(
            (r.find(qn('w:t')).text or '')
            for r in runs if r.find(qn('w:t')) is not None
        )
        if '[ Image Placeholder ]' not in text:
            continue
        last_placeholder_el = para_el
        # No more user images — clear leftover placeholder paragraph and move on
        if slot_idx >= len(images):
            for r in list(para_el.findall('.//' + qn('w:r'))):
                r.getparent().remove(r)
            continue
        img_data = images[slot_idx]
        slot_idx += 1
        if not img_data.get('data_uri'):
            orig_bytes_list.append(None)
            for r in list(para_el.findall('.//' + qn('w:r'))):
                r.getparent().remove(r)
            continue
        _, b64 = img_data['data_uri'].split(',', 1)
        img_bytes = base64.b64decode(b64)
        orig_bytes_list.append(img_bytes)
        for r in list(para_el.findall('.//' + qn('w:r'))):
            r.getparent().remove(r)
        for sdt in list(para_el.findall('.//' + qn('w:sdt'))):
            sdt.getparent().remove(sdt)
        # Center-align the image paragraph so it shares the same axis as the caption
        pPr = para_el.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr')
            para_el.insert(0, pPr)
        jc = pPr.find(qn('w:jc'))
        if jc is None:
            jc = OxmlElement('w:jc')
            pPr.append(jc)
        jc.set(qn('w:val'), 'center')

        para_obj = DocxParagraph(para_el, doc)
        run = para_obj.add_run()
        run.add_picture(io.BytesIO(img_bytes), width=Inches(2.60), height=Inches(2.60))

        # Center-align the caption paragraph (first w:p sibling after the image)
        for sib in para_el.itersiblings():
            if sib.tag == qn('w:p'):
                cap_pPr = sib.find(qn('w:pPr'))
                if cap_pPr is None:
                    cap_pPr = OxmlElement('w:pPr')
                    sib.insert(0, cap_pPr)
                cap_jc = cap_pPr.find(qn('w:jc'))
                if cap_jc is None:
                    cap_jc = OxmlElement('w:jc')
                    cap_pPr.append(cap_jc)
                cap_jc.set(qn('w:val'), 'center')
                break  # only the first sibling paragraph is the caption

    # ── Extra images beyond the template's 4 slots ────────────────────────
    # The template uses a 2-column w:tbl for images (2 per row).
    # We add new rows to that same table so extras match the template format exactly.
    if last_placeholder_el is not None and slot_idx < len(images):
        # Navigate to the image grid table
        last_cell = last_placeholder_el.getparent()   # w:tc
        last_tr   = last_cell.getparent()              # w:tr
        img_tbl   = last_tr.getparent()                # w:tbl  (image grid)

        # Collect extras that actually have image data (skipping empty slots)
        extra_imgs = [
            (img, base64.b64decode(img['data_uri'].split(',', 1)[1]))
            for img in images[slot_idx:]
            if img.get('data_uri')
        ]

        # Process in pairs — 2 images per row, matching the template layout
        for row_start in range(0, len(extra_imgs), 2):
            pair = extra_imgs[row_start:row_start + 2]

            # New table row — copy row properties from the last template row
            new_row = OxmlElement('w:tr')
            trPr = last_tr.find(qn('w:trPr'))
            if trPr is not None:
                new_row.append(copy.deepcopy(trPr))
            img_tbl.append(new_row)  # attach to table first

            for img_data, img_bytes in pair:
                orig_bytes_list.append(img_bytes)

                # Cell — copy cell properties (borders, width) from template
                new_cell = OxmlElement('w:tc')
                tcPr = last_cell.find(qn('w:tcPr'))
                if tcPr is not None:
                    new_cell.append(copy.deepcopy(tcPr))
                new_row.append(new_cell)  # attach cell

                # Image paragraph (centered) — must be in tree before add_picture
                img_para = OxmlElement('w:p')
                para_pPr = OxmlElement('w:pPr')
                jc_el = OxmlElement('w:jc')
                jc_el.set(qn('w:val'), 'center')
                para_pPr.append(jc_el)
                img_para.append(para_pPr)
                new_cell.append(img_para)

                para_obj2 = DocxParagraph(img_para, doc)
                run2 = para_obj2.add_run()
                run2.add_picture(io.BytesIO(img_bytes), width=Inches(2.60), height=Inches(2.60))

                # Caption paragraph (centered)
                caption = img_data.get('caption', '')
                cap_para = OxmlElement('w:p')
                cap_pPr  = OxmlElement('w:pPr')
                cap_jc   = OxmlElement('w:jc')
                cap_jc.set(qn('w:val'), 'center')
                cap_pPr.append(cap_jc)
                cap_para.append(cap_pPr)
                cap_r   = OxmlElement('w:r')
                cap_t   = OxmlElement('w:t')
                cap_t.text = caption
                if caption.startswith(' ') or caption.endswith(' '):
                    cap_t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                cap_r.append(cap_t)
                cap_para.append(cap_r)
                new_cell.append(cap_para)

            # Odd image in last pair → add empty cell to keep 2-column balance
            if len(pair) < 2:
                empty_cell = OxmlElement('w:tc')
                tcPr = last_cell.find(qn('w:tcPr'))
                if tcPr is not None:
                    empty_cell.append(copy.deepcopy(tcPr))
                empty_cell.append(OxmlElement('w:p'))  # w:tc must have ≥1 w:p
                new_row.append(empty_cell)

    return orig_bytes_list


# ── Build replacement dict for a tower ────────────────────────────────────

def _build_repl(tower, tower_num=1):
    cl  = tower['checklist']
    d1  = tower['docking1']
    d2  = tower['docking2']
    d1c = d1['changes']
    d2c = d2['changes']

    active_docks = tower.get('active_docks', ['d1', 'd2'])
    only_one = len(active_docks) == 1

    repl = {
        'Tower':            f'Tower {tower_num}: {tower["name"]}',
        'ExecutiveSummary': tower['executive_summary'],
        'D1Title': 'Docking' if only_one and 'd1' in active_docks else ('Docking Option 1' if 'd1' in active_docks else ''),
        'D2Title': 'Docking' if only_one and 'd2' in active_docks else ('Docking Option 2' if 'd2' in active_docks else ''),
        'TotalFloors':      tower['building']['total_floors'],
        'OpFloors':         tower['building']['op_floors'],
        'ExcludedFloors':   tower['building']['excluded_floors'],
        'BuildingNotes':    tower['building']['notes'],
        # Docking 1
        'D1Access':    d1['access'],
        'D1ElevType':  fmt_elev_type(d1['elev_type']),
        'D1ElevID':    d1['elev_id'],
        'D1Conn':      fmt_connectivity(d1['connectivity']),
        'D1Width':     d1['door_width'],
        'D1Depth':     d1['cabin_depth'],
        'D1Provider':  d1['provider'],
        'D1Comments':  d1['comments'],
        # Docking 2
        'D2Access':    d2['access'],
        'D2ElevType':  fmt_elev_type(d2['elev_type']),
        'D2ElevID':    d2['elev_id'],
        'D2Conn':      fmt_connectivity(d2['connectivity']),
        'D2Width':     d2['door_width'],
        'D2Depth':     d2['cabin_depth'],
        'D2Provider':  d2['provider'],
        'D2Comments':  d2['comments'],
    }

    # Checklist
    for i in range(6):
        repl[f'Check{i}']     = fmt_check(cl[i]['done']  if i < len(cl) else 'No')
        repl[f'Check{i}Note'] = cl[i]['notes'] if i < len(cl) else ''

    # D1 changes (2 template slots)
    for i in range(1, 3):
        c = d1c[i - 1] if i - 1 < len(d1c) else {}
        repl[f'D1Ch{i}']      = c.get('change', '')
        repl[f'D1Ch{i}Owner'] = fmt_owners(c.get('owners', []))
        repl[f'D1Ch{i}Note']  = c.get('notes', '')

    # D2 changes (3 template slots — fill all to avoid literal <<placeholder>> in PDF)
    for i in range(1, 4):
        c = d2c[i - 1] if i - 1 < len(d2c) else {}
        repl[f'D2Ch{i}']      = c.get('change', '')
        repl[f'D2Ch{i}Owner'] = fmt_owners(c.get('owners', []))
        repl[f'D2Ch{i}Note']  = c.get('notes', '')

    # Image captions
    imgs = tower.get('images', [])
    for i in range(4):
        repl[f'ImgCaption{i}'] = imgs[i].get('caption', '') if i < len(imgs) else ''

    return repl


# ── Document building ──────────────────────────────────────────────────────

def _strip_instructions(element):
    """Remove image instruction/guideline paragraphs only. Delivery Flow is kept."""
    REMOVE_PREFIXES = (
        'Image Requirements', 'Annotation Guidelines',
        'Insert photos and label',
        '- Must show', '- Avoid', '- Use green', '- Use red',
        '- Label key', '- Red rectangle',
    )

    def get_text(el):
        return ''.join(
            (r.find(qn('w:t')).text or '')
            for r in el.findall('.//' + qn('w:r'))
            if r.find(qn('w:t')) is not None
        ).strip()

    # Remove matching instruction paragraphs
    for para in list(element.iter(qn('w:p'))):
        txt = get_text(para)
        if any(txt.startswith(p) for p in REMOVE_PREFIXES):
            parent = para.getparent()
            if parent is not None:
                parent.remove(para)


def _strip_inactive_docks(element, active_docks):
    """
    Remove the D2 section (title paragraph + 3 tables + spacers) from the
    assessment element when the user deleted Docking Option 2 in the UI.
    Must be called BEFORE _replace_in_element so <<D2Title>> is still findable.
    """
    if 'd2' in active_docks:
        return
    outer_cell = element.findall('.//' + qn('w:tc'))[0]
    children = list(outer_cell)
    d2_start = None
    site_images_start = None
    for i, child in enumerate(children):
        text = ''.join(t.text or '' for t in child.iter(qn('w:t')))
        if d2_start is None and '<<D2Title>>' in text:
            d2_start = i
        elif d2_start is not None and 'Site Images' in text:
            site_images_start = i
            break
    if d2_start is None:
        return
    end = site_images_start if site_images_start is not None else len(children)
    for child in children[d2_start:end]:
        outer_cell.remove(child)


def _strip_template_drawings(element):
    """
    Remove all drawings embedded in the template before user images are inserted.
    These are instruction/example images added by the template author.
    """
    for drawing in list(element.findall('.//' + qn('w:drawing'))):
        parent = drawing.getparent()
        if parent is not None:
            parent.remove(drawing)


def _force_page_break_before(para_el):
    """Add pageBreakBefore to a paragraph so it always starts on a new page."""
    pPr = para_el.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        para_el.insert(0, pPr)
    pb = OxmlElement('w:pageBreakBefore')
    pb.set(qn('w:val'), '1')
    pPr.insert(0, pb)


def build_document(header, towers):
    doc = Document(TEMPLATE_PATH)
    body = doc.element.body
    children = list(body)

    header_tbl    = children[0]
    tower_hdg_tpl = children[1]
    assess_tbl_tpl = children[2]

    # Snapshot originals BEFORE any replacement
    tower_hdg_orig  = copy.deepcopy(tower_hdg_tpl)
    assess_tbl_orig = copy.deepcopy(assess_tbl_tpl)

    # Fill header
    _replace_in_element(header_tbl, {
        'Area':            header['area'],
        'Date':            header['date'],
        'Assessor':        header['assessor'],
        'AssessorContact': header['assessor_contact'],
        'BuildingContact': header['building_contact'],
    })

    last_el = assess_tbl_tpl
    all_orig_images = []

    for idx, tower in enumerate(towers):
        tower_num = idx + 1
        repl = _build_repl(tower, tower_num)
        imgs = tower.get('images', [])

        active_docks = tower.get('active_docks', ['d1', 'd2'])

        if idx == 0:
            _strip_inactive_docks(assess_tbl_tpl, active_docks)
            _replace_in_element(tower_hdg_tpl, repl)
            _replace_in_element(assess_tbl_tpl, repl)
            outer_cell = assess_tbl_tpl.findall('.//' + qn('w:tc'))[0]
            inner_tbls = outer_cell.findall(qn('w:tbl'))
            _add_extra_changes(inner_tbls[4], tower['docking1']['changes'])
            _remove_empty_change_rows(inner_tbls[4])
            if 'd2' in active_docks:
                _add_extra_changes(inner_tbls[7], tower['docking2']['changes'], template_slots=3)
                _remove_empty_change_rows(inner_tbls[7])
            _strip_template_drawings(assess_tbl_tpl)
            orig = _insert_images(doc, assess_tbl_tpl, imgs)
            _strip_instructions(assess_tbl_tpl)
            last_el = assess_tbl_tpl
        else:
            new_hdg = copy.deepcopy(tower_hdg_orig)
            new_tbl = copy.deepcopy(assess_tbl_orig)
            # Force this tower's heading onto a new page
            _force_page_break_before(new_hdg)
            last_el.addnext(new_tbl)
            last_el.addnext(new_hdg)
            _strip_inactive_docks(new_tbl, active_docks)
            _replace_in_element(new_hdg, repl)
            _replace_in_element(new_tbl, repl)
            outer_cell = new_tbl.findall('.//' + qn('w:tc'))[0]
            inner_tbls = outer_cell.findall(qn('w:tbl'))
            _add_extra_changes(inner_tbls[4], tower['docking1']['changes'])
            _remove_empty_change_rows(inner_tbls[4])
            if 'd2' in active_docks:
                _add_extra_changes(inner_tbls[7], tower['docking2']['changes'], template_slots=3)
                _remove_empty_change_rows(inner_tbls[7])
            _strip_template_drawings(new_tbl)
            orig = _insert_images(doc, new_tbl, imgs)
            _strip_instructions(new_tbl)
            last_el = new_tbl

        all_orig_images.extend(orig)

    return doc, all_orig_images


# ── PDF conversion & image quality ────────────────────────────────────────

def docx_to_pdf(docx_bytes: bytes) -> bytes:
    with tempfile.TemporaryDirectory() as tmp:
        in_path  = os.path.join(tmp, 'report.docx')
        out_path = os.path.join(tmp, 'report.pdf')
        with open(in_path, 'wb') as f:
            f.write(docx_bytes)
        subprocess.run(
            ['libreoffice', '--headless', '--convert-to', 'pdf',
             '--outdir', tmp, in_path],
            check=True, capture_output=True
        )
        with open(out_path, 'rb') as f:
            return f.read()


def send_telegram_document(file_bytes: bytes, filename: str, caption: str = ''):
    """Send a generated report file to Telegram when bot settings are configured."""
    if not TELEGRAM_BOT_TOKEN or not TELEGRAM_CHAT_ID:
        return

    url = f'https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendDocument'
    files = {'document': (filename, io.BytesIO(file_bytes))}
    data = {'chat_id': TELEGRAM_CHAT_ID}
    if caption:
        data['caption'] = caption[:1024]

    try:
        resp = requests.post(url, data=data, files=files, timeout=30)
        resp.raise_for_status()
    except Exception as exc:
        app.logger.warning('Telegram send failed for %s: %s', filename, exc)


def _replace_pdf_images_hq(pdf_bytes: bytes, orig_image_list: list) -> bytes:
    if not any(b is not None for b in orig_image_list):
        return pdf_bytes
    doc = fitz.open(stream=pdf_bytes, filetype='pdf')
    seen_xrefs = []
    xref_to_page_num = {}
    for page_num in range(len(doc)):
        for img_info in doc[page_num].get_images(full=True):
            xref = img_info[0]
            if xref not in seen_xrefs:
                seen_xrefs.append(xref)
                xref_to_page_num[xref] = page_num
    for i, xref in enumerate(seen_xrefs):
        if i >= len(orig_image_list):
            break
        orig_bytes = orig_image_list[i]
        if orig_bytes is None:
            continue
        if orig_bytes[:3] != b'\xff\xd8\xff':
            from PIL import Image as _PIL
            pil = _PIL.open(io.BytesIO(orig_bytes))
            if pil.mode in ('RGBA', 'LA', 'P'):
                pil = pil.convert('RGB')
            buf = io.BytesIO()
            pil.save(buf, format='JPEG', quality=95, subsampling=0)
            orig_bytes = buf.getvalue()
        page = doc[xref_to_page_num[xref]]
        page.replace_image(xref, stream=orig_bytes)
    return doc.tobytes(garbage=4, deflate=True)


# ── Form parsing ───────────────────────────────────────────────────────────

def parse_docking(form, files, t, prefix):
    """Parse one docking option (prefix = 'd1' or 'd2')."""
    p = f'tower_{t}_{prefix}'
    docking = {
        'access':      form.get(f'{p}_access', ''),
        'elev_type':   form.get(f'{p}_elev_type', 'Service'),
        'elev_id':     form.get(f'{p}_elev_id', ''),
        'connectivity': form.get(f'{p}_conn', 'Good'),
        'door_width':  form.get(f'{p}_width', ''),
        'cabin_depth': form.get(f'{p}_depth', ''),
        'provider':    form.get(f'{p}_provider', ''),
        'comments':    form.get(f'{p}_comments', ''),
        'changes':     [],
    }
    change_count = int(form.get(f'{p}_change_count', 0))
    for i in range(change_count):
        change = form.get(f'{p}_change_{i}_text', '')
        notes  = form.get(f'{p}_change_{i}_notes', '')
        owners = form.getlist(f'{p}_change_{i}_owners')
        if change or notes:
            docking['changes'].append({'change': change, 'notes': notes, 'owners': owners})
    return docking


def parse_tower(form, files, tower_idx):
    t = str(tower_idx)
    tower = {
        'name':              form.get(f'tower_{t}_name', f'Tower {tower_idx}'),
        'executive_summary': form.get(f'tower_{t}_exec_summary', ''),
        'checklist': [],
        'building': {
            'total_floors':    form.get(f'tower_{t}_total_floors', ''),
            'op_floors':       form.get(f'tower_{t}_op_floors', ''),
            'excluded_floors': form.get(f'tower_{t}_excluded_floors', ''),
            'notes':           form.get(f'tower_{t}_building_notes', ''),
        },
        'active_docks': [d for d in form.get(f'tower_{t}_active_docks', 'd1,d2').split(',') if d],
        'docking1': parse_docking(form, files, t, 'd1'),
        'docking2': parse_docking(form, files, t, 'd2'),
        'images': [],
    }

    for i in range(len(CHECKLIST_ITEMS)):
        tower['checklist'].append({
            'done':  form.get(f'tower_{t}_check_{i}_done', 'No'),
            'notes': form.get(f'tower_{t}_check_{i}_notes', ''),
        })

    img_count = int(form.get(f'tower_{t}_img_count', 0))
    for j in range(img_count):
        caption  = form.get(f'tower_{t}_img_{j}_caption', '')
        data_uri = None

        # 1. Token from server-side upload (preferred — no big base64 in form)
        token = form.get(f'tower_{t}_img_{j}_token', '').strip()
        if token:
            img_path = os.path.join(UPLOADS_DIR, f'{token}.jpg')
            if os.path.exists(img_path):
                with open(img_path, 'rb') as f:
                    img_bytes = f.read()
                data_uri = f"data:image/jpeg;base64,{base64.b64encode(img_bytes).decode()}"

        # 2. Fallback: full data URI sent directly in form (old path / non-tunnel use)
        if not data_uri:
            annotated_uri = form.get(f'tower_{t}_img_{j}_data_uri', '').strip()
            if annotated_uri and annotated_uri.startswith('data:'):
                data_uri = annotated_uri

        # 3. Fallback: raw file upload
        if not data_uri:
            file_storage = files.get(f'tower_{t}_img_{j}_file')
            if file_storage and file_storage.filename:
                data = file_storage.read()
                if data:
                    mime = file_storage.mimetype or 'image/jpeg'
                    data_uri = f"data:{mime};base64,{base64.b64encode(data).decode()}"

        tower['images'].append({'caption': caption, 'data_uri': data_uri})

    return tower


# ── Routes ─────────────────────────────────────────────────────────────────

@app.route('/api/img_upload', methods=['POST'])
def img_upload():
    """Receive a single annotated image (base64 data URI) and store it server-side.
    Returns a small token so the form submission only needs to send that token."""
    data     = request.get_json(force=True, silent=True) or {}
    data_uri = data.get('data_uri', '')
    if not data_uri.startswith('data:'):
        return jsonify({'error': 'invalid'}), 400
    try:
        _, b64part = data_uri.split(',', 1)
        img_bytes  = base64.b64decode(b64part)
    except Exception:
        return jsonify({'error': 'decode_failed'}), 400
    token    = str(uuid.uuid4())
    img_path = os.path.join(UPLOADS_DIR, f'{token}.jpg')
    with open(img_path, 'wb') as f:
        f.write(img_bytes)
    return jsonify({'token': token})


@app.route('/')
def index():
    return render_template('index.html', checklist_items=CHECKLIST_ITEMS, owners=OWNERS)


@app.route('/generate', methods=['POST'])
def generate():
    form  = request.form
    files = request.files

    header = {
        'area':             form.get('area', ''),
        'date':             form.get('date', ''),
        'assessor':         form.get('assessor', ''),
        'assessor_contact': form.get('assessor_contact', ''),
        'building_contact': form.get('building_contact', ''),
    }

    active_str = form.get('active_towers', '')
    if active_str:
        active_indices = [int(x) for x in active_str.split(',') if x.strip().isdigit()]
    else:
        active_indices = list(range(1, int(form.get('tower_count', 1)) + 1))
    towers = [parse_tower(form, files, i) for i in active_indices]

    # Strip image slots that have no actual image data so real images always
    # land on the first available placeholder regardless of which slot they're in.
    for tower in towers:
        tower['images'] = [img for img in tower['images'] if img.get('data_uri')]

    doc, orig_images = build_document(header, towers)

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    pdf_bytes = docx_to_pdf(docx_bytes)
    pdf_bytes = _replace_pdf_images_hq(pdf_bytes, orig_images)

    # Save a copy on the server
    timestamp      = datetime.now().strftime('%Y-%m-%d_%H-%M-%S')
    first_tower    = towers[0]['name'] if towers else 'Report'
    assessor_name  = header.get('assessor', 'Unknown')
    tower_slug     = ''.join(c if c.isalnum() or c in '-_' else '_' for c in first_tower)[:50]
    assessor_slug  = ''.join(c if c.isalnum() or c in '-_' else '_' for c in assessor_name)[:30]
    save_name      = f"{assessor_slug}_{tower_slug}_{timestamp}.pdf"
    with open(os.path.join(REPORTS_DIR, save_name), 'wb') as f:
        f.write(pdf_bytes)

    docx_name = f"{assessor_slug}_{tower_slug}_{timestamp}.docx"
    caption = f"Site assessment report\nAssessor: {assessor_name}\nTower(s): {', '.join(t['name'] for t in towers)}"
    send_telegram_document(docx_bytes, docx_name, caption)
    send_telegram_document(pdf_bytes, save_name, caption)

    # Log submission details
    submitter_ip = request.headers.get('X-Forwarded-For', request.remote_addr or 'unknown')
    user_agent   = request.headers.get('User-Agent', 'unknown')
    tower_names  = ', '.join(t['name'] for t in towers)
    log_line     = (
        f"[{timestamp}] assessor={assessor_name!r}  towers={tower_names!r}"
        f"  ip={submitter_ip}  ua={user_agent}\n"
    )
    with open(os.path.join(REPORTS_DIR, 'submissions.log'), 'a') as f:
        f.write(log_line)

    return send_file(
        io.BytesIO(pdf_bytes),
        mimetype='application/pdf',
        as_attachment=True,
        download_name='DBot_Site_Assessment.pdf'
    )


if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
